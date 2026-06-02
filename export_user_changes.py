"""
Export all inventory devices that a specific user has ever touched (added/edited/archived)
to an XLSX file in the standard inventory export format.

Usage:
    python export_user_changes.py dannhen
    python export_user_changes.py "Dan Hen"
"""

import sys
import os
import sqlite3
import openpyxl
from openpyxl.styles import Font

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE  = os.path.join(BASE_DIR, 'data', 'inventory.db')


def export_for_user(username: str):
    if not os.path.exists(DB_FILE):
        print(f"ERROR: Database not found at {DB_FILE}")
        sys.exit(1)

    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row

    # Find every device_id this user ever touched in the history log
    rows = conn.execute(
        'SELECT DISTINCT device_id FROM history WHERE LOWER(performed_by) = LOWER(?)',
        (username,)
    ).fetchall()

    device_ids = [r['device_id'] for r in rows if r['device_id'] is not None]

    if not device_ids:
        print(f"No history found for user '{username}'.")
        conn.close()
        sys.exit(0)

    print(f"Found {len(device_ids)} device(s) touched by '{username}'.")

    # Load the current state of those devices
    placeholders = ','.join('?' * len(device_ids))
    devices = conn.execute(
        f'SELECT * FROM inventory WHERE id IN ({placeholders})',
        device_ids
    ).fetchall()
    conn.close()

    # Build XLSX in the same format as the app's normal export
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{username} Changes"

    headers = [
        'ComputerName/ServiceTag', 'Location', 'Device Type', 'Model',
        'Username', 'Password', 'Island', 'Location Address w/zip',
        'Phone', 'Contact', 'Warranty End Date', 'Memory', 'CPU',
        'Operating Sys', 'NOTES:', 'Archived'
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for d in devices:
        d = dict(d)
        ws.append([
            d.get('serial_number', ''),
            d.get('location_code', ''),
            d.get('device_type', ''),
            d.get('model', ''),
            d.get('assigned_user', ''),
            d.get('password', ''),
            d.get('island', ''),
            '',                                        # Location Address w/zip
            d.get('phone', ''),
            '',                                        # Contact
            d.get('purchase_date', ''),
            '',                                        # Memory
            '',                                        # CPU
            '',                                        # Operating Sys
            d.get('notes', ''),
            'Yes' if d.get('archived') else 'No',
        ])

    # Auto-fit column widths (same as app export)
    for col in ws.columns:
        max_len = max((len(str(cell.value)) if cell.value else 0) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    safe_name = username.replace(' ', '_').lower()
    out_path = os.path.join(BASE_DIR, f'inventory_{safe_name}_changes.xlsx')
    wb.save(out_path)
    print(f"Exported {len(devices)} device(s) to {out_path}")


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python export_user_changes.py <username>")
        sys.exit(1)
    export_for_user(sys.argv[1])
